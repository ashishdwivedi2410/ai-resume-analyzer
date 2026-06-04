import io
from pathlib import Path
from typing import Tuple
from fastapi import UploadFile, HTTPException, status
from config import get_settings
from utils.logger import app_logger

settings = get_settings()


async def parse_resume(file: UploadFile) -> Tuple[str, str]:
    """
    Parse uploaded resume file and extract raw text.
    Returns: (raw_text, file_type)
    """
    filename = file.filename or ""
    ext = Path(filename).suffix.lower()

    if ext not in [".pdf", ".docx"]:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file type '{ext}'. Only PDF and DOCX are allowed."
        )

    content = await file.read()
    file_size_mb = len(content) / (1024 * 1024)

    if file_size_mb > settings.max_file_size_mb:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File too large ({file_size_mb:.1f}MB). Max allowed: {settings.max_file_size_mb}MB"
        )

    app_logger.info(f"Parsing resume: {filename} ({file_size_mb:.2f}MB)")

    if ext == ".pdf":
        raw_text = _extract_pdf_text(content)
    else:
        raw_text = _extract_docx_text(content)

    if not raw_text.strip():
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Could not extract text from the uploaded file. It may be scanned or image-based."
        )

    app_logger.info(f"Extracted {len(raw_text)} characters from {filename}")
    return raw_text, ext.lstrip(".")


def _extract_pdf_text(content: bytes) -> str:
    """Extract text from PDF using PyMuPDF."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=content, filetype="pdf")
        text_parts = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text_parts.append(page.get_text("text"))
        doc.close()
        return "\n".join(text_parts)
    except ImportError:
        app_logger.error("PyMuPDF (fitz) not installed")
        raise HTTPException(status_code=500, detail="PDF parsing library not available")
    except Exception as e:
        app_logger.error(f"PDF extraction error: {e}")
        raise HTTPException(status_code=422, detail=f"Failed to parse PDF: {str(e)}")


def _extract_docx_text(content: bytes) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(content))
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())

        return "\n".join(paragraphs)
    except ImportError:
        app_logger.error("python-docx not installed")
        raise HTTPException(status_code=500, detail="DOCX parsing library not available")
    except Exception as e:
        app_logger.error(f"DOCX extraction error: {e}")
        raise HTTPException(status_code=422, detail=f"Failed to parse DOCX: {str(e)}")
    